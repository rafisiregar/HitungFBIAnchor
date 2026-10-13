"""
fbi_calc.py — Perhitungan Fee Based Income (FBI) merchant EDC.
Logika IDENTIK dengan js/fbi.js (versi web). Dipakai untuk batch / verifikasi / GitHub Actions.

Pakai sebagai CLI:
    python fbi_calc.py RAW.xlsx MASTER.xlsx -o HASIL.xlsx [--join TID|MID|GROUP|NAMA_MERCHANT]
                                                          [--qris MAPPING|OFF_US|ON_US] [--bulan Juni]
Pakai sebagai modul:
    from fbi_calc import FBICalculator
    calc = FBICalculator(join_key="TID")      # <- ubah di sini kalau mau join pakai GROUP, dll.
    df_hasil, rekap = calc.compute(df_raw, df_master)
"""
from __future__ import annotations
import argparse, re, sys
from dataclasses import dataclass, field
import pandas as pd

# =====================================================================
#  KONFIGURASI — satu-satunya bagian yang perlu diubah bila aturan berubah
# =====================================================================
JSON_KET_MAPPING = [
    {"KATEGORI": "BLANK",          "GROUPING": "KREDIT", "KLASIFIKASI": "ON US",  "KET GABUNGAN": "KREDIT ON US",  "KET MDR": "MDR On Us"},
    {"KATEGORI": "BRIZZI",         "GROUPING": "BRIZZI", "KLASIFIKASI": "ON US",  "KET GABUNGAN": "BRIZZI ON US",  "KET MDR": "MDR Uang Elektronik"},
    {"KATEGORI": "KB-JCB",         "GROUPING": "KREDIT", "KLASIFIKASI": "OFF US", "KET GABUNGAN": "KREDIT OFF US", "KET MDR": "MDR Off Us", "KET INTERCHANGE": "ITC JCB"},
    {"KATEGORI": "KB-LAINNYA",     "GROUPING": "KREDIT", "KLASIFIKASI": "OFF US", "KET GABUNGAN": "KREDIT OFF US", "KET MDR": "MDR Off Us", "KET INTERCHANGE": "ITC LAINNYA"},
    {"KATEGORI": "KB-LAINNYA-NPG", "GROUPING": "DEBIT",  "KLASIFIKASI": "OFF US", "KET GABUNGAN": "DEBIT OFF US",  "KET MDR": "MDR NPG"},
    {"KATEGORI": "KB-MC",          "GROUPING": "KREDIT", "KLASIFIKASI": "OFF US", "KET GABUNGAN": "KREDIT OFF US", "KET MDR": "MDR Off Us", "KET INTERCHANGE": "ITC MC"},
    {"KATEGORI": "KB-MC-NPG",      "GROUPING": "DEBIT",  "KLASIFIKASI": "OFF US", "KET GABUNGAN": "DEBIT OFF US",  "KET MDR": "MDR NPG"},
    {"KATEGORI": "KB-VISA",        "GROUPING": "KREDIT", "KLASIFIKASI": "OFF US", "KET GABUNGAN": "KREDIT OFF US", "KET MDR": "MDR Off Us", "KET INTERCHANGE": "ITC VISA"},
    {"KATEGORI": "KB-VISA-NPG",    "GROUPING": "DEBIT",  "KLASIFIKASI": "OFF US", "KET GABUNGAN": "DEBIT OFF US",  "KET MDR": "MDR NPG"},
    {"KATEGORI": "KD-BRI",         "GROUPING": "DEBIT",  "KLASIFIKASI": "ON US",  "KET GABUNGAN": "DEBIT ON US",   "KET MDR": "MDR Debit"},
    {"KATEGORI": "KD-BRI-TARTUN",  "GROUPING": "DEBIT",  "KLASIFIKASI": "ON US",  "KET GABUNGAN": "DEBIT ON US",   "KET MDR": "MDR Debit"},
    {"KATEGORI": "KK-BRI-JCB",     "GROUPING": "KREDIT", "KLASIFIKASI": "ON US",  "KET GABUNGAN": "KREDIT ON US",  "KET MDR": "MDR On Us"},
    {"KATEGORI": "KK-BRI-MC",      "GROUPING": "KREDIT", "KLASIFIKASI": "ON US",  "KET GABUNGAN": "KREDIT ON US",  "KET MDR": "MDR On Us"},
    {"KATEGORI": "KK-BRI-VISA",    "GROUPING": "KREDIT", "KLASIFIKASI": "ON US",  "KET GABUNGAN": "KREDIT ON US",  "KET MDR": "MDR On Us"},
    {"KATEGORI": "KK-LAINNYA",     "GROUPING": "KREDIT", "KLASIFIKASI": "OFF US", "KET GABUNGAN": "KREDIT OFF US", "KET MDR": "MDR Off Us", "KET INTERCHANGE": "ITC LAINNYA"},
    # QRIS: KLASIFIKASI "QRIS" dibaca engine => ada KET INTERCHANGE = OFF US, tidak ada = ON US
    {"KATEGORI": "QRIS",           "GROUPING": "QRIS",   "KLASIFIKASI": "QRIS",   "KET GABUNGAN": "QRIS", "KET MDR": "MDR QRIS On Us"},
    {"KATEGORI": "QRIS-CB",        "GROUPING": "QRIS",   "KLASIFIKASI": "QRIS",   "KET GABUNGAN": "QRIS", "KET MDR": "MDR QRIS Off US", "KET INTERCHANGE": "QRIS"},
    {"KATEGORI": "QRIS-NFC",       "GROUPING": "QRIS",   "KLASIFIKASI": "QRIS",   "KET GABUNGAN": "QRIS", "KET MDR": "MDR QRIS Off US", "KET INTERCHANGE": "QRIS"},
    {"KATEGORI": "QRIS-TITO",      "GROUPING": "QRIS",   "KLASIFIKASI": "QRIS",   "KET GABUNGAN": "QRIS", "KET MDR": "MDR QRIS Off US", "KET INTERCHANGE": "QRIS"},
    {"KATEGORI": "QRIS ON US",     "GROUPING": "QRIS",   "KLASIFIKASI": "QRIS",   "KET GABUNGAN": "QRIS", "KET MDR": "MDR QRIS On US"},
    {"KATEGORI": "QRIS OFF US",    "GROUPING": "QRIS",   "KLASIFIKASI": "QRIS",   "KET GABUNGAN": "QRIS", "KET MDR": "MDR QRIS Off US", "KET INTERCHANGE": "QRIS"},
]

JOIN_KEYS        = ["TID", "MID", "GROUP", "NAMA_MERCHANT"]
JOIN_KEY_DEFAULT = "TID"
QRIS_MODES       = ["MAPPING", "OFF_US", "ON_US"]     # MAPPING = ikuti JSON_KET_MAPPING; OFF_US/ON_US = override semua QRIS raw
QRIS_MODE_DEFAULT = "MAPPING"
SV_SUFFIX        = " SV"                               # "KB-MC SV" = "KB-MC" + " SV"
NETT_COLUMN_BY_GROUPING = {"DEBIT": "DEBIT", "QRIS": "QRIS"}   # pengali nett OFF US tanpa KET INTERCHANGE
OUTPUT_COLUMNS   = ["BRIZZI", "CC ON US", "CC OFF US", "DEBIT ON US", "DEBIT OFF US", "QRIS ON US", "QRIS OFF US"]
GABUNGAN_TO_OUTPUT = {
    "BRIZZI ON US": "BRIZZI", "KREDIT ON US": "CC ON US", "KREDIT OFF US": "CC OFF US",
    "DEBIT ON US": "DEBIT ON US", "DEBIT OFF US": "DEBIT OFF US", "QRIS ON US": "QRIS ON US", "QRIS OFF US": "QRIS OFF US",
}
BLOCK_LABELS = {"GROSS": "Gross FBI {BULAN}", "ITC": "Interchange Fee {BULAN}", "NET": "Net FBI {BULAN}",
                "TOTAL": "Total Net FBI {BULAN}", "KET": "KETERANGAN"}
MASTER_SHEET_NAME, MASTER_HEADER_ROW, RAW_HEADER_ROW = "MDR DAN NETT", 2, 2
UNMATCHED_LABEL = "MDR TIDAK DITEMUKAN"


# =====================================================================
#  ENGINE
# =====================================================================
def norm(s) -> str:
    return re.sub(r"\s+", " ", str("" if s is None else s)).strip().upper()


def num(v) -> float:
    try:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return 0.0
        if isinstance(v, (int, float)):
            return float(v)
        return float(re.sub(r"[^\d.\-eE]", "", str(v)) or 0)
    except ValueError:
        return 0.0


@dataclass
class FBICalculator:
    join_key: str = JOIN_KEY_DEFAULT
    qris_mode: str = QRIS_MODE_DEFAULT
    mapping: list = field(default_factory=lambda: [dict(m) for m in JSON_KET_MAPPING])

    # ---- mapping efektif (identik dengan effectiveMapping di fbi.js)
    def effective_mapping(self) -> list[dict]:
        m = [dict(r) for r in self.mapping]
        if self.qris_mode in ("OFF_US", "ON_US"):
            ref_name = "QRIS OFF US" if self.qris_mode == "OFF_US" else "QRIS ON US"
            ref = next((r for r in m if norm(r["KATEGORI"]) == ref_name), None)
            if ref:
                for r in m:
                    if norm(r["GROUPING"]) == "QRIS":
                        r["KET MDR"] = ref["KET MDR"]
                        if ref.get("KET INTERCHANGE"): r["KET INTERCHANGE"] = ref["KET INTERCHANGE"]
                        else: r.pop("KET INTERCHANGE", None)
        for r in m:
            kls, gab = norm(r["KLASIFIKASI"]), norm(r["KET GABUNGAN"])
            if kls == "QRIS": kls = "OFF US" if r.get("KET INTERCHANGE") else "ON US"
            if gab == "QRIS": gab = "QRIS OFF US" if kls == "OFF US" else "QRIS ON US"
            r["_KLASIFIKASI"], r["_GABUNGAN"] = kls, gab
        return m

    @staticmethod
    def _colmap(df: pd.DataFrame) -> dict:
        return {norm(c): c for c in df.columns}

    def compute(self, df_raw: pd.DataFrame, df_master: pd.DataFrame, bulan: str = "", rekap_col: str = "GROUP"):
        """Mengembalikan (df_hasil, df_rekap). df_hasil = raw + 21 kolom FBI + Total + KETERANGAN."""
        mapping = self.effective_mapping()
        rc, mc = self._colmap(df_raw), self._colmap(df_master)
        jk = norm(self.join_key)
        if jk not in rc: raise KeyError(f"Kolom join '{self.join_key}' tidak ada di raw")
        if jk not in mc: raise KeyError(f"Kolom join '{self.join_key}' tidak ada di master")

        master_idx = {}
        for _, r in df_master.iterrows():
            k = norm(r[mc[jk]])
            if k and k not in master_idx: master_idx[k] = r

        L = lambda t: t.replace("{BULAN}", bulan or "")
        blocks = [("GROSS", "gross"), ("ITC", "itc"), ("NET", "net")]
        out_rows = []
        for _, raw in df_raw.iterrows():
            master = master_idx.get(norm(raw[rc[jk]]))
            res = {b: {c: 0.0 for c in OUTPUT_COLUMNS} for _, b in blocks}
            warn = []
            if master is None:
                ket = UNMATCHED_LABEL
            else:
                for m in mapping:
                    col = rc.get(norm(m["KATEGORI"] + SV_SUFFIX))
                    sv = num(raw[col]) if col else 0.0
                    if not sv: continue
                    out_col = GABUNGAN_TO_OUTPUT.get(m["_GABUNGAN"])
                    if not out_col or not m.get("KET MDR"): continue
                    mdr_col = mc.get(norm(m["KET MDR"]))
                    if mdr_col is None: warn.append(f"KOLOM {m['KET MDR']} TIDAK ADA")
                    gross = sv * (num(master[mdr_col]) if mdr_col else 0.0)
                    nett = 0.0
                    if m["_KLASIFIKASI"] == "OFF US":
                        nett_name = m.get("KET INTERCHANGE") or NETT_COLUMN_BY_GROUPING.get(norm(m["GROUPING"]), "")
                        nett_col = mc.get(norm(nett_name)) if nett_name else None
                        if nett_name and nett_col is None: warn.append(f"KOLOM {nett_name} TIDAK ADA")
                        nett = num(master[nett_col]) if nett_col else 0.0
                    itc = sv * nett
                    res["gross"][out_col] += gross; res["itc"][out_col] += itc; res["net"][out_col] += gross - itc
                ket = "; ".join(dict.fromkeys(warn))
            line = {}
            for lab, key in blocks:
                for c in OUTPUT_COLUMNS: line[(L(BLOCK_LABELS[lab]), c)] = res[key][c]
            total = sum(res["net"].values())
            line[(L(BLOCK_LABELS["TOTAL"]), L(BLOCK_LABELS["TOTAL"]))] = total
            line[(BLOCK_LABELS["KET"], BLOCK_LABELS["KET"])] = ket or None
            out_rows.append(line)

        df_calc = pd.DataFrame(out_rows, index=df_raw.index)
        df_calc.columns = pd.MultiIndex.from_tuples(df_calc.columns)
        df_hasil = pd.concat([df_raw, df_calc], axis=1)

        # rekap
        rk = rc.get(norm(rekap_col))
        net_cols = [(L(BLOCK_LABELS["NET"]), c) for c in OUTPUT_COLUMNS]
        tot_col = (L(BLOCK_LABELS["TOTAL"]), L(BLOCK_LABELS["TOTAL"]))
        g = df_raw[rk].fillna("(kosong)") if rk else pd.Series(["(semua)"] * len(df_raw), index=df_raw.index)
        total_label = f"TOTAL NET FBI {bulan}".strip()
        rekap = pd.DataFrame({
            rekap_col: g, "JML TID": 1,
            "GROSS FBI": df_calc[L(BLOCK_LABELS["GROSS"])].sum(axis=1),
            "INTERCHANGE FEE": df_calc[L(BLOCK_LABELS["ITC"])].sum(axis=1),
            **{f"NET {c}": df_calc[nc] for c, nc in zip(OUTPUT_COLUMNS, net_cols)},
            total_label: df_calc[tot_col],
        }).groupby(rekap_col, as_index=False).sum().sort_values(total_label, ascending=False)
        return df_hasil, rekap


# =====================================================================
#  I/O helper
# =====================================================================
def read_sheet(path: str, header_row: int, sheet=None) -> tuple[pd.DataFrame, str]:
    """Baca sheet dengan header di baris ke-N; kembalikan (df, label_baris_1)."""
    xl = pd.ExcelFile(path)
    if sheet is None:
        sheet = next((s for s in xl.sheet_names if re.search(r"mentah|raw|data", s, re.I)), xl.sheet_names[0])
    elif sheet not in xl.sheet_names:
        sheet = next((s for s in xl.sheet_names if norm(s) == norm(sheet)), xl.sheet_names[0])
    top = pd.read_excel(path, sheet_name=sheet, header=None, nrows=header_row - 1)
    title = next((str(v).strip() for v in top.values.ravel() if pd.notna(v) and str(v).strip()), "") if len(top) else ""
    df = pd.read_excel(path, sheet_name=sheet, header=header_row - 1)
    df = df.loc[:, [c for c in df.columns if not str(c).startswith("Unnamed")]].dropna(how="all")
    df.columns = [str(c).strip() for c in df.columns]
    return df, title


def write_result(path: str, df_hasil: pd.DataFrame, rekap: pd.DataFrame, calc: FBICalculator, bulan: str, raw_title: str = ""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(); ws = wb.active; ws.title = "HASIL PERHITUNGAN"
    n_raw = sum(1 for c in df_hasil.columns if not isinstance(c, tuple))
    h1, h2 = [], []
    for c in df_hasil.columns:
        if isinstance(c, tuple): h1.append(c[0]); h2.append(c[1])
        else: h1.append(None); h2.append(c)
    if raw_title: h1[0] = raw_title
    ws.append(h1); ws.append(h2)
    for row in df_hasil.itertuples(index=False): ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    # merge header blok
    start = None
    for i, v in enumerate(h1 + [None]):
        if start is not None and (v != h1[start]):
            if i - 1 > start: ws.merge_cells(start_row=1, start_column=start + 1, end_row=1, end_column=i)
            start = None
        if v is not None and start is None: start = i
    for i, (a, b) in enumerate(zip(h1, h2)):
        if a is not None and a == b: ws.merge_cells(start_row=1, start_column=i + 1, end_row=2, end_column=i + 1)
    fill = {"Gross": "DDEBF7", "Inter": "FCE4D6", "Net F": "E2EFDA", "Total": "FFF2CC", "KETER": "EDEDED"}
    for i, (a, b) in enumerate(zip(h1, h2), start=1):
        for r in (1, 2):
            cell = ws.cell(row=r, column=i); cell.font = Font(bold=True); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            key = str(a or b)[:5]
            if i > n_raw and key in fill: cell.fill = PatternFill("solid", fgColor=fill[key])
        ws.column_dimensions[get_column_letter(i)].width = 16 if i > n_raw else min(max(len(str(b)) + 2, 8), 30)
        if i > n_raw:
            for r in range(3, ws.max_row + 1): ws.cell(row=r, column=i).number_format = "#,##0.00;[Red]-#,##0.00"
    ws.freeze_panes = ws.cell(row=3, column=4)

    ws2 = wb.create_sheet("REKAP")
    ws2.append(list(rekap.columns))
    for row in rekap.itertuples(index=False): ws2.append(list(row))
    ws2.append(["TOTAL"] + [rekap[c].sum() for c in rekap.columns[1:]])
    for c in ws2[1]: c.font = Font(bold=True)
    for c in ws2[ws2.max_row]: c.font = Font(bold=True)
    for col in range(3, ws2.max_column + 1):
        for r in range(2, ws2.max_row + 1): ws2.cell(row=r, column=col).number_format = "#,##0;[Red]-#,##0"
    ws2.column_dimensions["A"].width = 28
    for col in range(2, ws2.max_column + 1): ws2.column_dimensions[get_column_letter(col)].width = 16

    ws3 = wb.create_sheet("PARAMETER")
    ws3.append(["Parameter", "Nilai"]); ws3.append(["Bulan", bulan]); ws3.append(["Kunci join", calc.join_key]); ws3.append(["Mode QRIS raw", calc.qris_mode])
    ws3.append(["Jumlah TID", len(df_hasil)]); ws3.append([])
    ws3.append(["KATEGORI", "GROUPING", "KLASIFIKASI", "KLASIFIKASI EFEKTIF", "KOLOM OUTPUT", "KET MDR", "KET INTERCHANGE / NETT"])
    for m in calc.effective_mapping():
        ws3.append([m["KATEGORI"], m["GROUPING"], m["KLASIFIKASI"], m["_KLASIFIKASI"], GABUNGAN_TO_OUTPUT.get(m["_GABUNGAN"], "(tidak dihitung)"),
                    m.get("KET MDR", ""), m.get("KET INTERCHANGE") or (NETT_COLUMN_BY_GROUPING.get(norm(m["GROUPING"]), "") if m["_KLASIFIKASI"] == "OFF US" else "")])
    for col, w in zip("ABCDEFG", (18, 12, 12, 18, 16, 22, 24)): ws3.column_dimensions[col].width = w
    wb.save(path)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Hitung FBI merchant dari raw SV + master MDR")
    ap.add_argument("raw"); ap.add_argument("master"); ap.add_argument("-o", "--output", default="FBI_hasil.xlsx")
    ap.add_argument("--join", default=JOIN_KEY_DEFAULT, choices=JOIN_KEYS); ap.add_argument("--qris", default=QRIS_MODE_DEFAULT, choices=QRIS_MODES)
    ap.add_argument("--bulan", default=None, help="label bulan di header (default: baris 1 raw)")
    ap.add_argument("--raw-sheet", default=None); ap.add_argument("--master-sheet", default=MASTER_SHEET_NAME); ap.add_argument("--rekap", default="GROUP")
    a = ap.parse_args(argv)
    df_raw, title = read_sheet(a.raw, RAW_HEADER_ROW, a.raw_sheet)
    df_master, _ = read_sheet(a.master, MASTER_HEADER_ROW, a.master_sheet)
    bulan = a.bulan if a.bulan is not None else title
    calc = FBICalculator(join_key=a.join, qris_mode=a.qris)
    df_hasil, rekap = calc.compute(df_raw, df_master, bulan=bulan, rekap_col=a.rekap)
    write_result(a.output, df_hasil, rekap, calc, bulan, raw_title=title)
    tot = df_hasil[[c for c in df_hasil.columns if isinstance(c, tuple) and c[0].startswith("Total")][0]].sum()
    unmatched = int((df_hasil[("KETERANGAN", "KETERANGAN")] == UNMATCHED_LABEL).sum())
    print(f"OK: {len(df_hasil)} TID, {unmatched} tanpa MDR, Total Net FBI {bulan}: {tot:,.2f} -> {a.output}")


if __name__ == "__main__":
    sys.exit(main())
